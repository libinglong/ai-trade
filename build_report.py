# -*- coding: utf-8 -*-
"""
build_report.py
===============
读取 7-9 月成交记录 Excel，按「品种」识别每一笔「建仓 -> 清仓」往返交易，
并生成一个自包含、可交互的前端页面 report.html。

口径：
  - 品种：合约代码去掉末尾月份数字（jd2609 -> JD，CF609 -> CF）。
  - 一笔交易：某品种的品种级净持仓从 0 建仓，到最后清仓（手数回到 0）。
  - 移仓：同一品种内不同合约间的切换（近月平仓 + 远月建仓），品种级
    净持仓未回到 0，因此不算清仓，仍属同一笔交易。

用法：
    python build_report.py "2026 7-9月成交.xlsx"
    python build_report.py            # 不传参时自动取当前目录下第一个 .xlsx
"""

import sys
import re
import datetime as dt
from collections import defaultdict
from html import escape
from pathlib import Path

import openpyxl

# 品种代码 -> 中文名称（可自行增补）
VARIETY_NAMES = {
    'CF': '棉花', 'MA': '甲醇', 'PK': '花生', 'PR': '瓶片', 'SF': '硅铁', 'TA': 'PTA',
    'AL': '铝', 'AO': '氧化铝', 'BU': '沥青', 'CU': '铜', 'EB': '苯乙烯',
    'EG': '乙二醇', 'FU': '燃料油', 'JD': '鸡蛋', 'JM': '焦煤', 'LC': '碳酸锂',
    'LH': '生猪', 'LU': '低硫燃料油', 'M': '豆粕', 'NR': '20号胶', 'P': '棕榈油',
    'PB': '铅', 'PG': '液化石油气', 'PP': '聚丙烯', 'PS': '聚苯乙烯', 'RU': '天然橡胶',
    'SN': '锡', 'V': '聚氯乙烯', 'ZN': '锌',
}

CLOSE_MARKS = ('平', '平昨', '平今')


def extract_variety(contract):
    return re.sub(r'\d+$', '', str(contract).strip()).upper()


def _fnum(v):
    if v is None:
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _parse_date(s):
    s = str(s).strip()
    m = re.fullmatch(r'(\d{4})(\d{2})(\d{2})', s)
    return dt.date(*map(int, m.groups())) if m else None


def _canon_price(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v).strip()
    return f'{f:.3f}'.rstrip('0').rstrip('.')


def _fmt_lots(v):
    return str(int(round(v))) if abs(v - round(v)) < 1e-9 else f'{v:.2f}'


def _fmt_money(v):
    return '0.00' if abs(v) < 1e-9 else f'{v:+,.2f}'


def aggregate(excel_path):
    """解析 Excel，返回 (varieties, totals)。

    varieties: [{code, name, ntrades, open, close, pl, trades:[...]}]
    每笔 trade: {open_date, close_date, days, contracts, open_lots,
                close_lots, pl, groups:[{price, side, kp, count, lots, pl}]}
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    header = [str(ws.cell(1, c).value).strip() for c in range(1, ws.max_column + 1)]
    col = {name: i for i, name in enumerate(header, 1)}

    state = defaultdict(lambda: {
        'open': 0.0, 'close': 0.0, 'pl': 0.0,
        'trades': [], 'cur': None, 'pos': 0.0,
    })

    current_date = None
    for r in range(2, ws.max_row + 1):
        dv = ws.cell(r, col['日期']).value
        if dv is not None and str(dv).strip() != '' and not str(dv).startswith('合计'):
            current_date = _parse_date(dv)

        contract = ws.cell(r, col['合约']).value
        if contract is None or str(contract).strip() == '':
            continue
        kp = str(ws.cell(r, col['开平']).value or '').strip()
        if kp == '':
            continue

        variety = extract_variety(contract)
        st = state[variety]
        lots = _fnum(ws.cell(r, col['手数']).value)
        side = str(ws.cell(r, col['买卖']).value).strip()
        price = str(ws.cell(r, col['成交价']).value or '').strip()
        pl = _fnum(ws.cell(r, col['平仓盈亏']).value)

        before = st['pos']
        st['pos'] += lots if side == '买' else -lots
        after = st['pos']

        if kp == '开':
            st['open'] += lots
        elif kp in CLOSE_MARKS:
            st['close'] += lots
            st['pl'] += pl

        # 交易边界：0 -> 非0 开新一笔；非0 -> 0 完成一笔
        if before == 0 and after != 0:
            st['cur'] = {'open_date': current_date, 'fills': [], 'contracts': set(),
                         'open_lots': 0.0, 'close_lots': 0.0, 'pl': 0.0}
        if st['cur'] is not None:
            cur = st['cur']
            cur['fills'].append({'date': current_date, 'contract': str(contract).strip(),
                                 'side': side, 'kp': kp, 'lots': lots,
                                 'price': price, 'pl': pl})
            cur['contracts'].add(str(contract).strip())
            if kp == '开':
                cur['open_lots'] += lots
            elif kp in CLOSE_MARKS:
                cur['close_lots'] += lots
                cur['pl'] += pl
        if before != 0 and after == 0:
            cur = st['cur']
            st['trades'].append(_finalize_trade(cur, current_date))
            st['cur'] = None

    varieties = []
    totals = {'nvarieties': 0, 'ntrades': 0, 'open': 0.0, 'close': 0.0, 'pl': 0.0}
    for code in sorted(state):
        st = state[code]
        if st['trades']:
            varieties.append({
                'code': code,
                'name': VARIETY_NAMES.get(code, code),
                'ntrades': len(st['trades']),
                'open': st['open'],
                'close': st['close'],
                'pl': st['pl'],
                'trades': sorted(st['trades'], key=lambda t: t['open_date'] or ''),
            })
            totals['nvarieties'] += 1
            totals['ntrades'] += len(st['trades'])
            totals['open'] += st['open']
            totals['close'] += st['close']
            totals['pl'] += st['pl']

    varieties.sort(key=lambda v: -abs(v['pl']))
    return varieties, totals


def _finalize_trade(cur, close_date):
    groups = []
    index = {}
    for f in cur['fills']:
        price = _canon_price(f['price'])
        key = (price, f['side'], f['kp'])
        if key not in index:
            index[key] = {'price': price, 'side': f['side'], 'kp': f['kp'],
                          'count': 0, 'lots': 0.0, 'pl': 0.0}
            groups.append(index[key])
        g = index[key]
        g['count'] += 1
        g['lots'] += f['lots']
        g['pl'] += f['pl']
    open_date = cur['open_date']
    days = (close_date - open_date).days if (open_date and close_date) else 0
    return {
        'open_date': open_date.strftime('%Y-%m-%d') if open_date else '',
        'close_date': close_date.strftime('%Y-%m-%d') if close_date else '',
        'days': days,
        'contracts': sorted(cur['contracts']),
        'open_lots': cur['open_lots'],
        'close_lots': cur['close_lots'],
        'pl': cur['pl'],
        'groups': groups,
    }


# ----------------------------------------------------------------------
# 渲染 HTML
# ----------------------------------------------------------------------

CSS = """
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:-apple-system,"Segoe UI","Microsoft YaHei",sans-serif;background:#f5f6f8;color:#1f2430;padding:24px 16px;line-height:1.5}
.wrap{max-width:1100px;margin:0 auto}
header{background:#fff;border:1px solid #e6e8ec;border-radius:12px;padding:20px 24px;margin-bottom:16px}
h1{font-size:20px;font-weight:700}
.sub{color:#8a8f99;font-size:13px;margin-top:4px}
.stats{display:flex;gap:12px;margin-top:16px;flex-wrap:wrap}
.stat{background:#f7f8fa;border:1px solid #eceef2;border-radius:10px;padding:10px 18px;min-width:120px}
.stat .k{font-size:12px;color:#8a8f99}
.stat .v{font-size:22px;font-weight:700;margin-top:2px;font-variant-numeric:tabular-nums}
.pos{color:#d9403a}.neg{color:#0a8f4f}
.card{background:#fff;border:1px solid #e6e8ec;border-radius:12px;margin-bottom:16px;overflow:hidden}
.card>h2{font-size:15px;padding:14px 20px;border-bottom:1px solid #eef0f3}
table{width:100%;border-collapse:collapse;font-size:14px}
th,td{padding:9px 12px;text-align:left;border-bottom:1px solid #eef0f3;vertical-align:middle}
th{background:#fafbfc;color:#6a7080;font-weight:600;font-size:12px}
td.num,th.num{text-align:right;font-variant-numeric:tabular-nums}
tr.total td{background:#fafbfc;font-weight:700;border-top:1px solid #e6e8ec}
.variety-block{border-bottom:1px solid #eef0f3}
.variety-block:last-child{border-bottom:none}
.variety-head{display:flex;align-items:center;gap:10px;padding:12px 20px;background:#fafbfc;border-bottom:1px solid #eef0f3;flex-wrap:wrap}
.variety-head .vname{font-weight:700;font-size:14px}
.variety-head .vcode{color:#8a8f99;font-size:12px}
.variety-head .vmeta{margin-left:auto;font-size:13px;color:#6a7080;display:flex;gap:16px;flex-wrap:wrap}
details{border-bottom:1px solid #eef0f3}
details:last-child{border-bottom:none}
summary{cursor:pointer;list-style:none;padding:10px 20px;font-size:13px;display:flex;align-items:center;gap:12px;flex-wrap:wrap}
summary:hover{background:#fafbfc}
summary::-webkit-details-marker{display:none}
.caret{color:#a0a5b0;transition:transform .15s;display:inline-block}
details[open]>.caret-wrap .caret{transform:rotate(90deg)}
details[open]>summary{background:#f7f8fa}
.caret-wrap{display:inline-flex;align-items:center}
.trade-date{font-weight:600}
.trade-contracts{display:inline-flex;gap:4px;flex-wrap:wrap}
.contract-tag{background:#eef1f6;border-radius:6px;padding:1px 7px;font-size:11px;color:#4a5265}
.trade-meta{color:#6a7080}
.trade-pl{margin-left:auto;font-weight:700}
.detail{padding:2px 20px 16px}
.detail-head{background:#fafbfc;border:1px solid #eef0f3;border-radius:8px;padding:12px 14px;margin:4px 0 12px;font-size:13px;display:flex;flex-wrap:wrap;gap:6px 24px}
.detail-head b{font-weight:700}
.detail-table td,.detail-table th{font-size:13px}
.badge{display:inline-block;border-radius:5px;padding:0 6px;font-size:11px;line-height:18px}
.badge.open{background:#e8f3ff;color:#1668b3}
.badge.close{background:#fff3e6;color:#c26a00}
.badge.buy{background:#fdeaea;color:#d9403a}
.badge.sell{background:#e9f7ef;color:#0a8f4f}
.foot{text-align:center;color:#a0a5b0;font-size:12px;padding:16px}
"""


def _render_stats(totals):
    pl_cls = 'pos' if totals['pl'] >= 0 else 'neg'
    return (
        '<div class="stats">'
        f'<div class="stat"><div class="k">品种数</div><div class="v">{totals["nvarieties"]}</div></div>'
        f'<div class="stat"><div class="k">交易笔数</div><div class="v">{totals["ntrades"]}</div></div>'
        f'<div class="stat"><div class="k">开仓手数</div><div class="v">{_fmt_lots(totals["open"])}</div></div>'
        f'<div class="stat"><div class="k">平仓手数</div><div class="v">{_fmt_lots(totals["close"])}</div></div>'
        f'<div class="stat"><div class="k">总盈亏</div><div class="v {pl_cls}">{_fmt_money(totals["pl"])}</div></div>'
        '</div>'
    )


def _render_summary_table(varieties, totals):
    rows = []
    for v in varieties:
        pl_cls = 'pos' if v['pl'] >= 0 else 'neg'
        rows.append(
            f'<tr><td>{escape(v["code"])}</td><td>{escape(v["name"])}</td>'
            f'<td class="num">{v["ntrades"]}</td>'
            f'<td class="num">{_fmt_lots(v["open"])}</td>'
            f'<td class="num">{_fmt_lots(v["close"])}</td>'
            f'<td class="num {pl_cls}">{_fmt_money(v["pl"])}</td></tr>'
        )
    total_pl_cls = 'pos' if totals['pl'] >= 0 else 'neg'
    total = (
        '<tr class="total"><td>合计</td><td></td>'
        f'<td class="num">{totals["ntrades"]}</td>'
        f'<td class="num">{_fmt_lots(totals["open"])}</td>'
        f'<td class="num">{_fmt_lots(totals["close"])}</td>'
        f'<td class="num {total_pl_cls}">{_fmt_money(totals["pl"])}</td></tr>'
    )
    return (
        '<div class="card"><h2>品种汇总</h2><table>'
        '<thead><tr><th>品种</th><th>品种名称</th><th class="num">交易笔数</th>'
        '<th class="num">开仓手数</th><th class="num">平仓手数</th><th class="num">盈亏</th></tr></thead>'
        '<tbody>' + ''.join(rows) + total + '</tbody></table></div>'
    )


def _render_trade(trade, trade_no):
    pl_cls = 'pos' if trade['pl'] >= 0 else 'neg'
    contracts = ''.join(
        f'<span class="contract-tag">{escape(c)}</span>' for c in trade['contracts'])
    span = f'{trade["open_date"]} → {trade["close_date"]}'
    if trade['days'] == 0:
        span += '（当日）'
    else:
        span += f'（{trade["days"]} 天）'

    summary = (
        f'<span class="caret-wrap"><span class="caret">▶</span></span>'
        f'<span class="trade-no">#{trade_no}</span>'
        f'<span class="trade-date">{span}</span>'
        f'<span class="trade-contracts">{contracts}</span>'
        f'<span class="trade-meta">开 {_fmt_lots(trade["open_lots"])} 手 · '
        f'平 {_fmt_lots(trade["close_lots"])} 手</span>'
        f'<span class="trade-pl {pl_cls}">{_fmt_money(trade["pl"])}</span>'
    )

    head = (
        '<div class="detail-head">'
        f'<span><b>开仓</b> {trade["open_date"]}</span>'
        f'<span><b>平仓</b> {trade["close_date"]}</span>'
        f'<span><b>时间跨度</b> {trade["days"]} 天</span>'
        f'<span><b>开仓手数</b> {_fmt_lots(trade["open_lots"])}</span>'
        f'<span><b>平仓手数</b> {_fmt_lots(trade["close_lots"])}</span>'
        f'<span><b>总盈亏</b> <span class="{pl_cls}">{_fmt_money(trade["pl"])}</span></span>'
        '</div>'
    )

    body_rows = []
    for g in trade['groups']:
        kp_badge = 'open' if g['kp'] == '开' else 'close'
        side_badge = 'buy' if g['side'] == '买' else 'sell'
        pl_text = '—' if g['kp'] == '开' else _fmt_money(g['pl'])
        pl_cls = '' if g['kp'] == '开' else ('pos' if g['pl'] >= 0 else 'neg')
        body_rows.append(
            f'<tr><td class="num">{escape(g["price"])}</td>'
            f'<td><span class="badge {side_badge}">{g["side"]}</span></td>'
            f'<td><span class="badge {kp_badge}">{g["kp"]}</span></td>'
            f'<td class="num">{g["count"]}</td>'
            f'<td class="num">{_fmt_lots(g["lots"])}</td>'
            f'<td class="num {pl_cls}">{pl_text}</td></tr>'
        )
    table = (
        '<table class="detail-table"><thead><tr>'
        '<th class="num">点位(成交价)</th><th>买卖</th><th>开平</th>'
        '<th class="num">成交笔数</th><th class="num">聚合手数</th><th class="num">盈亏</th>'
        '</tr></thead><tbody>' + ''.join(body_rows) + '</tbody></table>'
    )

    return (
        '<details><summary>' + summary + '</summary>'
        '<div class="detail">' + head + table + '</div></details>'
    )


def _render_trades(varieties):
    blocks = []
    trade_no = 0
    for v in varieties:
        pl_cls = 'pos' if v['pl'] >= 0 else 'neg'
        trades_html = ''.join(
            _render_trade(t, i + 1) for i, t in enumerate(v['trades']))
        blocks.append(
            '<div class="variety-block"><div class="variety-head">'
            f'<span class="vname">{escape(v["name"])}</span>'
            f'<span class="vcode">{escape(v["code"])}</span>'
            '<span class="vmeta">'
            f'<span>交易 {v["ntrades"]} 笔</span>'
            f'<span>开 {_fmt_lots(v["open"])} 手</span>'
            f'<span>平 {_fmt_lots(v["close"])} 手</span>'
            f'<span class="{pl_cls}">盈亏 {_fmt_money(v["pl"])}</span>'
            '</span></div>' + trades_html + '</div>'
        )
    return '<div class="card"><h2>交易明细（点击查看详情）</h2>' + ''.join(blocks) + '</div>'


def render_html(varieties, totals, src_name):
    return (
        '<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        '<title>成交记录品种汇总</title><style>' + CSS + '</style></head><body>'
        '<div class="wrap">'
        '<header><h1>成交记录品种汇总</h1>'
        f'<div class="sub">数据来源：{escape(src_name)} · 按品种识别「建仓 → 清仓」往返交易</div>'
        + _render_stats(totals) + '</header>'
        + _render_summary_table(varieties, totals)
        + _render_trades(varieties)
        + '<div class="foot">由 build_report.py 生成</div>'
        '</div></body></html>'
    )


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    if argv is None:
        argv = sys.argv[1:]

    if argv:
        path = argv[0]
    else:
        cands = [p for p in Path('.').glob('*.xlsx') if not p.name.startswith('~$')]
        if not cands:
            print('未找到 Excel 文件。用法：python build_report.py "成交记录.xlsx"')
            return 1
        path = str(cands[0])

    varieties, totals = aggregate(path)
    out_path = Path(path).with_name('report.html')
    out_path.write_text(render_html(varieties, totals, Path(path).name), encoding='utf-8')
    print(f'品种 {totals["nvarieties"]} 个，交易 {totals["ntrades"]} 笔，'
          f'总盈亏 {_fmt_money(totals["pl"])}')
    print(f'已生成前端页面：{out_path}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
