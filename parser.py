# -*- coding: utf-8 -*-
"""
parser.py
=========
把「成交记录」Excel 解析成结构化的成交明细（fills）列表。

Excel 列（成交记录）：
  日期 | 交易所 | 合约 | 买卖 | 开平 | 手数 | 成交价 | 成交额 | 平仓盈亏 | 手续费 | 投保 | 权利金收支

本模块只负责「读 Excel → 结构化数据」，不掺杂业务聚合逻辑。
"""

import re
import datetime as dt

import openpyxl

# 「开平」字段中属于平仓的取值
CLOSE_MARKS = ('平', '平昨', '平今')


def extract_variety(contract):
    """从合约代码里提取品种代码：jd2609 -> JD，CF609 -> CF，v2701 -> V。"""
    return re.sub(r'\d+$', '', str(contract).strip()).upper()


def _fnum(v):
    if v is None:
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _parse_date(s):
    """把 '20260702' 这类日期字符串解析成 date；解析失败返回 None。"""
    s = str(s).strip()
    m = re.fullmatch(r'(\d{4})(\d{2})(\d{2})', s)
    return dt.date(*map(int, m.groups())) if m else None


def _canon_price(v):
    """把成交价规整成可比较的字符串（用于「相同点位」合并）。"""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v).strip()
    return f'{f:.3f}'.rstrip('0').rstrip('.')


def parse_excel(excel_path):
    """解析 Excel，返回结构化成交明细列表。

    每条明细（fill）字段：
      date      datetime.date  交易日
      contract  str            合约代码（如 CF609）
      variety   str            品种代码（如 CF）
      side      str            买卖：买 / 卖
      kp        str            开平：开 / 平 / 平今 / 平昨
      lots      float          手数
      price     str            成交价（规整后的字符串）
      amount    float          成交额（= 成交价 × 合约乘数 × 手数，即货值）
      pl        float          平仓盈亏
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    header = [str(ws.cell(1, c).value or '').strip() for c in range(1, ws.max_column + 1)]
    col = {name: i for i, name in enumerate(header, 1)}

    fills = []
    current_date = None
    for r in range(2, ws.max_row + 1):
        dv = ws.cell(r, col['日期']).value
        if dv is not None and str(dv).strip() != '' and not str(dv).startswith('合计'):
            current_date = _parse_date(dv)

        contract = ws.cell(r, col['合约']).value
        if contract is None or str(contract).strip() == '':
            continue
        kp = str(ws.cell(r, col['开平']).value or '').strip()
        if kp == '':
            continue

        fills.append({
            'date': current_date,
            'contract': str(contract).strip(),
            'variety': extract_variety(contract),
            'side': str(ws.cell(r, col['买卖']).value or '').strip(),
            'kp': kp,
            'lots': _fnum(ws.cell(r, col['手数']).value),
            'price': _canon_price(ws.cell(r, col['成交价']).value),
            'amount': _fnum(ws.cell(r, col['成交额']).value),
            'pl': _fnum(ws.cell(r, col['平仓盈亏']).value),
        })

    return fills
